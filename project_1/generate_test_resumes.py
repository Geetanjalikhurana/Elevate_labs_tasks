"""
generate_test_resumes.py
Generates 5 realistic sample test resumes (PDF & DOCX formats) for testing Smart Resume Parser.
"""

import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx


TEST_RESUMES_DIR = "test_resumes"
os.makedirs(TEST_RESUMES_DIR, exist_ok=True)

# Sample Resumes Data
SAMPLE_RESUMES = [
    {
        "filename": "1_Alex_Rivers_Data_Scientist.pdf",
        "name": "Alex Rivers",
        "email": "alex.rivers@email.com",
        "phone": "+1 (555) 234-5678",
        "linkedin": "linkedin.com/in/alex-rivers-data",
        "github": "github.com/alexrivers",
        "summary": "Results-driven Senior Data Scientist with 6+ years of experience in Machine Learning, Deep Learning, NLP, and Predictive Analytics. Proven track record of deploying robust ML models in production.",
        "skills": "Python, R, PyTorch, TensorFlow, Scikit-Learn, Pandas, NumPy, NLP, spaCy, HuggingFace, SQL, PostgreSQL, Docker, AWS, Streamlit, Git",
        "education": "Master of Science in Computer Science - Stanford University (2018 - 2020)\nBachelor of Technology in Computer Engineering - MIT (2014 - 2018)",
        "experience": "Senior Data Scientist | AI Solutions Inc. (2021 - Present)\n- Built LLM-powered document parsing and NLP extraction pipelines improving search efficiency by 45%.\n- Fine-tuned BERT and spaCy transformers for automated entity extraction.\n\nData Scientist | TechCorp Analytics (2018 - 2021)\n- Developed predictive machine learning models using XGBoost and Random Forest with 92% accuracy.\n- Designed interactive dashboards with Streamlit and Plotly for executive stakeholders."
    },
    {
        "filename": "2_Priya_Sharma_FullStack_Developer.pdf",
        "name": "Priya Sharma",
        "email": "priya.sharma@techdev.org",
        "phone": "+91 98765 43210",
        "linkedin": "linkedin.com/in/priyasharma-dev",
        "github": "github.com/priyasharma99",
        "summary": "Versatile Full Stack Developer with 4 years of experience crafting scalable web applications using React, Node.js, Python, and cloud infrastructure.",
        "skills": "JavaScript, TypeScript, Python, React.js, Next.js, Node.js, Express.js, FastAPI, HTML5, CSS3, TailwindCSS, MongoDB, PostgreSQL, REST API, GraphQL, Docker, Git",
        "education": "Bachelor of Engineering in Information Technology - Delhi Technological University (2018 - 2022)",
        "experience": "Full Stack Engineer | CloudScale Tech (2022 - Present)\n- Architected RESTful microservices with Node.js and Express, handling over 1M daily requests.\n- Developed high-performance frontend interfaces using React.js and TypeScript.\n\nSoftware Engineering Intern | WebMatrix (2021 - 2022)\n- Integrated MongoDB database schemas and implemented JWT authentication protocols."
    },
    {
        "filename": "3_Marcus_Vance_DevOps_Architect.docx",
        "name": "Marcus Vance",
        "email": "marcus.vance@cloudops.io",
        "phone": "+1 415 889 9012",
        "linkedin": "linkedin.com/in/marcusvance-devops",
        "github": "github.com/mvance-ops",
        "summary": "Lead DevOps & Cloud Infrastructure Architect with 8+ years designing automated CI/CD pipelines, Kubernetes orchestration, and AWS enterprise environments.",
        "skills": "AWS, Azure, Docker, Kubernetes, Terraform, Ansible, Jenkins, GitHub Actions, Linux, Bash, Python, Prometheus, Grafana, Nginx, CI/CD, System Design",
        "education": "Bachelor of Science in Information Systems - University of Washington (2014 - 2018)",
        "experience": "Lead Cloud Infrastructure Architect | Enterprise Systems (2020 - Present)\n- Managed multi-region Kubernetes clusters on AWS EKS serving 10M+ active users.\n- Automated infrastructure provisioning with Terraform, cutting setup time by 70%.\n\nSenior DevOps Engineer | InfraSolutions (2018 - 2020)\n- Implemented GitOps CI/CD pipelines using GitHub Actions and Helm charts."
    },
    {
        "filename": "4_Sarah_Jenkins_Frontend_Engineer.pdf",
        "name": "Sarah Jenkins",
        "email": "sarah.jenkins@designcode.com",
        "phone": "+1 650 334 1122",
        "linkedin": "linkedin.com/in/sarahjenkins-fe",
        "github": "github.com/sjenkins-code",
        "summary": "Creative Frontend Engineer with 5 years specializing in modern UI/UX design systems, Web performance optimization, React, Vue.js, and TypeScript.",
        "skills": "JavaScript, TypeScript, React.js, Vue.js, Redux, HTML5, CSS3, Sass, TailwindCSS, Webpack, Vite, Figma, Jest, Cypress, UI/UX Design",
        "education": "Bachelor of Fine Arts in Web Design & Interactive Media - RISD (2017 - 2021)",
        "experience": "Senior Frontend Developer | Digital Craft Studio (2021 - Present)\n- Led frontend engineering team in building responsive React applications with 99% Lighthouse score.\n- Created reusable UI component library in Storybook with Figma synchronization."
    },
    {
        "filename": "5_Rahul_Verma_Python_Backend_Developer.docx",
        "name": "Rahul Verma",
        "email": "rahul.verma@pybackend.net",
        "phone": "+91 91234 56789",
        "linkedin": "linkedin.com/in/rahulverma-py",
        "github": "github.com/rverma-backend",
        "summary": "Dedicated Python Backend Developer with 3 years experience building high-throughput REST APIs, asynchronous microservices, and database optimization.",
        "skills": "Python, Django, Flask, FastAPI, SQL, MySQL, Redis, Celery, RESTful API, PyTest, Git, Docker, Linux, Microservices",
        "education": "Bachelor of Technology in Computer Science - BITS Pilani (2019 - 2023)",
        "experience": "Backend Engineer | DataFlow Tech (2023 - Present)\n- Designed asynchronous REST APIs with FastAPI and Pydantic for high-speed data ingestion.\n- Optimized complex MySQL SQL queries and implemented Redis caching, reducing query latency by 50%.\n\nBackend Developer Intern | PyLabs Solutions (2022 - 2023)\n- Developed unit and integration tests using pytest achieving 95% test coverage."
    }
]


def create_pdf_resume(data, filepath):
    doc = SimpleDocTemplate(filepath, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=20, leading=24, textColor=colors.HexColor('#1E293B'))
    contact_style = ParagraphStyle('Contact', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor('#475569'))
    heading_style = ParagraphStyle('Heading', parent=styles['Heading2'], fontSize=12, leading=16, textColor=colors.HexColor('#2563EB'), spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor('#334155'))
    
    elements = []
    
    # Name & Contact
    elements.append(Paragraph(f"<b>{data['name']}</b>", title_style))
    contact_info = f"Email: {data['email']} | Phone: {data['phone']} | LinkedIn: {data['linkedin']} | GitHub: {data['github']}"
    elements.append(Paragraph(contact_info, contact_style))
    elements.append(Spacer(1, 8))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=10))
    
    # Summary
    elements.append(Paragraph("<b>PROFESSIONAL SUMMARY</b>", heading_style))
    elements.append(Paragraph(data['summary'], body_style))
    elements.append(Spacer(1, 6))
    
    # Skills
    elements.append(Paragraph("<b>TECHNICAL SKILLS</b>", heading_style))
    elements.append(Paragraph(data['skills'], body_style))
    elements.append(Spacer(1, 6))
    
    # Experience
    elements.append(Paragraph("<b>WORK EXPERIENCE</b>", heading_style))
    for line in data['experience'].split('\n'):
        if line.strip():
            elements.append(Paragraph(line.replace('\n', '<br/>'), body_style))
    elements.append(Spacer(1, 6))
    
    # Education
    elements.append(Paragraph("<b>EDUCATION</b>", heading_style))
    for line in data['education'].split('\n'):
        if line.strip():
            elements.append(Paragraph(line, body_style))

    doc.build(elements)


def create_docx_resume(data, filepath):
    doc = docx.Document()
    
    # Name
    p_title = doc.add_paragraph()
    run_title = p_title.add_run(data['name'])
    run_title.bold = True
    run_title.font.size = docx.shared.Pt(20)
    
    # Contact
    p_contact = doc.add_paragraph(f"Email: {data['email']} | Phone: {data['phone']}\nLinkedIn: {data['linkedin']} | GitHub: {data['github']}")
    
    # Summary
    h_summary = doc.add_heading('PROFESSIONAL SUMMARY', level=2)
    doc.add_paragraph(data['summary'])
    
    # Skills
    h_skills = doc.add_heading('TECHNICAL SKILLS', level=2)
    doc.add_paragraph(data['skills'])
    
    # Experience
    h_exp = doc.add_heading('WORK EXPERIENCE', level=2)
    doc.add_paragraph(data['experience'])
    
    # Education
    h_edu = doc.add_heading('EDUCATION', level=2)
    doc.add_paragraph(data['education'])

    doc.save(filepath)


def generate_all_resumes():
    print("Generating sample test resumes...")
    generated_paths = []
    for data in SAMPLE_RESUMES:
        filepath = os.path.join(TEST_RESUMES_DIR, data['filename'])
        if data['filename'].endswith('.pdf'):
            create_pdf_resume(data, filepath)
        else:
            create_docx_resume(data, filepath)
        print(f"Generated: {filepath}")
        generated_paths.append(filepath)
    return generated_paths


if __name__ == "__main__":
    generate_all_resumes()
