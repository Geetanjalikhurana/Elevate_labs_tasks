"""
resume_parser.py
Core Resume Parsing Engine using PyMuPDF / pymupdf, python-docx, spaCy, and Regex.
"""

import re
import os
import io
import pymupdf
import docx
import spacy
from skills_db import get_flat_skill_map, SKILL_TAXONOMY

# Load spaCy model
try:
    nlp = spacy.load("en_core_web_sm")
except Exception:
    import spacy.cli
    spacy.cli.download("en_core_web_sm")
    nlp = spacy.load("en_core_web_sm")


class ResumeParser:
    def __init__(self):
        self.skill_map = get_flat_skill_map()

    def extract_text_from_pdf(self, file_source):
        """Extract text from PDF file path or bytes."""
        text = ""
        try:
            if isinstance(file_source, (str, bytes)):
                if isinstance(file_source, bytes):
                    doc = pymupdf.open(stream=file_source, filetype="pdf")
                else:
                    doc = pymupdf.open(file_source)
                
                for page in doc:
                    text += page.get_text("text") + "\n"
                doc.close()
            else:
                doc = pymupdf.open(stream=file_source.read(), filetype="pdf")
                for page in doc:
                    text += page.get_text("text") + "\n"
                doc.close()
        except Exception as e:
            text = f"Error extracting text from PDF: {str(e)}"
        return text

    def extract_text_from_docx(self, file_source):
        """Extract text from DOCX file path or bytes/file-like object."""
        text = ""
        try:
            if isinstance(file_source, bytes):
                doc_file = io.BytesIO(file_source)
            elif hasattr(file_source, 'read'):
                doc_file = io.BytesIO(file_source.read())
            else:
                doc_file = file_source

            doc = docx.Document(doc_file)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            full_text.append(cell.text.strip())
            
            text = "\n".join(full_text)
        except Exception as e:
            text = f"Error extracting text from DOCX: {str(e)}"
        return text

    def clean_text(self, text):
        """Clean and normalize extracted text."""
        if not text:
            return ""
        text = re.sub(r'[\r\t]', ' ', text)
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        return "\n".join(lines)

    def extract_email(self, text):
        """Extract email address using regex."""
        email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
        match = re.search(email_pattern, text)
        return match.group(0) if match else "N/A"

    def extract_phone(self, text):
        """Extract phone number using regex."""
        phone_patterns = [
            r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}',
            r'\+?\d{1,3}[-.\s]?\d{5}[-.\s]?\d{5}',
            r'\+?\d{1,4}[-.\s]?\d{9,10}',
            r'\b\d{10}\b'
        ]
        for pattern in phone_patterns:
            match = re.search(pattern, text)
            if match:
                phone = match.group(0).strip()
                if len(re.sub(r'\D', '', phone)) >= 10:
                    return phone
        return "N/A"

    def extract_links(self, text):
        """Extract LinkedIn and GitHub profile links."""
        linkedin = re.search(r'(?:https?://)?(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+/?', text, re.IGNORECASE)
        github = re.search(r'(?:https?://)?(?:www\.)?github\.com/[a-zA-Z0-9_-]+/?', text, re.IGNORECASE)
        
        return {
            "LinkedIn": linkedin.group(0) if linkedin else "N/A",
            "GitHub": github.group(0) if github else "N/A"
        }

    def extract_name(self, text):
        """Extract candidate name using spaCy NER and header heuristics."""
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        if not lines:
            return "N/A"

        header_text = "\n".join(lines[:5])
        doc = nlp(header_text)
        for ent in doc.ents:
            if ent.label_ == "PERSON":
                name = ent.text.strip()
                name_lower = name.lower()
                # Exclude if matched string is a skill or common heading
                if name_lower not in self.skill_map and not re.search(r'\d|email|phone|http|summary|experience|education|skills', name_lower):
                    if len(name.split()) <= 4 and name_lower not in ["curriculum vitae", "resume", "cv"]:
                        return name

        # Heuristic fallback: First clean line that looks like a name
        for line in lines[:3]:
            line_lower = line.lower()
            if line_lower not in self.skill_map and not re.search(r'\d|@|http|www|email|phone|summary|experience|education|skills', line_lower):
                if len(line.split()) in [2, 3, 4] and line_lower not in ["curriculum vitae", "resume", "cv"]:
                    return line.title()

        return lines[0].title() if lines else "Candidate Name"

    def extract_skills(self, text):
        """
        Extract skills based on dictionary taxonomy and regex boundary matching.
        """
        text_lower = text.lower()
        extracted_skills = {}
        matched_skills_set = set()

        for canonical_lower, info in self.skill_map.items():
            escaped = re.escape(canonical_lower)
            pattern = r'(?:\b|(?<=\W))' + escaped + r'(?:\b|(?=\W))'
            if re.search(pattern, text_lower):
                matched_skills_set.add(info["display_name"])
                cat = info["category"]
                if cat not in extracted_skills:
                    extracted_skills[cat] = []
                if info["display_name"] not in extracted_skills[cat]:
                    extracted_skills[cat].append(info["display_name"])

        all_skills_list = sorted(list(matched_skills_set))
        return {
            "categorized_skills": extracted_skills,
            "all_skills": all_skills_list,
            "total_skills_count": len(all_skills_list)
        }

    def extract_education(self, text):
        """Extract education details (degrees, field of study, passing year, GPA)."""
        education_list = []
        
        degree_patterns = [
            r'\b(B\.?Tech|Bachelor\s+of\s+Technology|B\.?E\.?|Bachelor\s+of\s+Engineering|B\.?S\.?|Bachelor\s+of\s+Science|B\.?C\.?A\.?)\b',
            r'\b(M\.?Tech|Master\s+of\s+Technology|M\.?E\.?|Master\s+of\s+Engineering|M\.?S\.?|Master\s+of\s+Science|M\.?C\.?A\.?|M\.?B\.?A\.?)\b',
            r'\b(Ph\.?D\.?|Doctor\s+of\s+Philosophy)\b',
            r'\b(Diploma|High\s+School|Secondary\s+School)\b'
        ]

        lines = text.split('\n')
        in_edu_section = False

        for line in lines:
            line_str = line.strip()
            if re.search(r'\b(EDUCATION|ACADEMICS|ACADEMIC BACKGROUND|QUALIFICATIONS)\b', line_str, re.IGNORECASE):
                in_edu_section = True
                continue

            if in_edu_section:
                if re.search(r'\b(EXPERIENCE|WORK|PROJECTS|SKILLS|CERTIFICATIONS|SUMMARY)\b', line_str, re.IGNORECASE) and not re.search(r'\b(EDUCATION)\b', line_str, re.IGNORECASE):
                    in_edu_section = False
                    break
                
                if line_str:
                    education_list.append(line_str)
            else:
                for pat in degree_patterns:
                    if re.search(pat, line_str, re.IGNORECASE) and line_str not in education_list:
                        education_list.append(line_str)
                        break

        if not education_list:
            for line in lines:
                for pat in degree_patterns:
                    if re.search(pat.strip(), line.strip(), re.IGNORECASE):
                        education_list.append(line.strip())

        return education_list[:5]

    def extract_experience(self, text):
        """Extract work experience entries and calculate estimated total experience in years."""
        exp_lines = []
        lines = text.split('\n')
        in_exp_section = False

        for line in lines:
            line_str = line.strip()
            if re.search(r'\b(WORK EXPERIENCE|EXPERIENCE|EMPLOYMENT|WORK HISTORY|PROFESSIONAL EXPERIENCE)\b', line_str, re.IGNORECASE):
                in_exp_section = True
                continue
            
            if in_exp_section:
                if re.search(r'\b(EDUCATION|SKILLS|PROJECTS|CERTIFICATIONS|ACHIEVEMENTS)\b', line_str, re.IGNORECASE):
                    in_exp_section = False
                    break
                if line_str:
                    exp_lines.append(line_str)

        years = re.findall(r'\b(19\d{2}|20\d{2})\b', text)
        years = [int(y) for y in years if 1980 <= int(y) <= 2030]
        
        estimated_years = 0
        if years:
            min_yr, max_yr = min(years), max(years)
            current_yr = 2026
            if max_yr == min_yr:
                estimated_years = 1
            else:
                estimated_years = min(max_yr, current_yr) - min_yr

        return {
            "details": exp_lines[:10],
            "estimated_years": estimated_years
        }

    def calculate_ats_score(self, parsed_resume, job_description):
        """
        Calculate ATS compatibility score between candidate resume and a target job description.
        """
        if not job_description or not job_description.strip():
            return {
                "ats_score": 0,
                "matched_skills": [],
                "missing_skills": [],
                "match_percentage": "0%",
                "recommendation": "Provide a job description to calculate ATS Match Score."
            }

        jd_lower = job_description.lower()
        
        jd_skills = set()
        for canonical_lower, info in self.skill_map.items():
            escaped = re.escape(canonical_lower)
            pattern = r'(?:\b|(?<=\W))' + escaped + r'(?:\b|(?=\W))'
            if re.search(pattern, jd_lower):
                jd_skills.add(info["display_name"])

        candidate_skills = set(parsed_resume.get("skills", {}).get("all_skills", []))

        if not jd_skills:
            jd_words = set(re.findall(r'\b\w{4,}\b', jd_lower))
            res_words = set(re.findall(r'\b\w{4,}\b', parsed_resume.get("raw_text", "").lower()))
            overlap = jd_words.intersection(res_words)
            score = round((len(overlap) / max(len(jd_words), 1)) * 100, 1) if jd_words else 50.0
            return {
                "ats_score": min(score, 100.0),
                "matched_skills": list(overlap)[:10],
                "missing_skills": list(jd_words - res_words)[:10],
                "match_percentage": f"{min(score, 100.0)}%",
                "recommendation": "Moderate keyword match found."
            }

        matched = candidate_skills.intersection(jd_skills)
        missing = jd_skills - candidate_skills

        score = round((len(matched) / len(jd_skills)) * 100, 1)

        if score >= 85:
            recommendation = "Excellent Match! High candidate suitability for this position."
        elif score >= 60:
            recommendation = "Good Match. Consider adding missing keywords to boost alignment."
        else:
            missing_sample = ", ".join(list(missing)[:5])
            recommendation = f"Low Match. Missing key requirements: {missing_sample}."

        return {
            "ats_score": score,
            "matched_skills": sorted(list(matched)),
            "missing_skills": sorted(list(missing)),
            "match_percentage": f"{score}%",
            "recommendation": recommendation
        }

    def parse(self, file_source, filename="resume.pdf"):
        """Main parsing method for a given resume file."""
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".docx":
            raw_text = self.extract_text_from_docx(file_source)
        else:
            raw_text = self.extract_text_from_pdf(file_source)

        cleaned_text = self.clean_text(raw_text)
        
        email = self.extract_email(cleaned_text)
        phone = self.extract_phone(cleaned_text)
        links = self.extract_links(cleaned_text)
        name = self.extract_name(cleaned_text)
        skills = self.extract_skills(cleaned_text)
        education = self.extract_education(cleaned_text)
        experience = self.extract_experience(cleaned_text)

        return {
            "filename": filename,
            "name": name,
            "email": email,
            "phone": phone,
            "links": links,
            "skills": skills,
            "education": education,
            "experience": experience,
            "raw_text": cleaned_text
        }
